#!/usr/bin/env python3
"""
Build pitch/Shotorko-Presenter-Brief.docx — everything to know before presenting.

The deck carries keywords; this document carries the understanding. It is generated, not
hand-typed, so it cannot drift: the on-slide text is read back out of the built .pptx and the
sentences to say are read out of pitch/deck-and-qa.md. If a slide changes, re-run this.

    python3 dev/build-brief.py
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

ROOT = Path(__file__).resolve().parent.parent
PPTX = ROOT / "pitch" / "Shotorko-GP-FutureMakers-2026.pptx"
MD = ROOT / "pitch" / "deck-and-qa.md"
OUT = ROOT / "pitch" / "Shotorko-Presenter-Brief.docx"

AMBER = RGBColor(0xB2, 0x6A, 0x00)
INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x5A, 0x5A, 0x5A)

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(5)
st.paragraph_format.line_spacing = 1.12

for s in doc.sections:
    s.left_margin = s.right_margin = Inches(0.75)
    s.top_margin = s.bottom_margin = Inches(0.7)


def h1(text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.color.rgb = AMBER
        r.font.size = Pt(16)
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    return p


def h2(text):
    p = doc.add_heading(text, level=2)
    for r in p.runs:
        r.font.color.rgb = INK
        r.font.size = Pt(12)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    return p


def para(text, size=10.5, color=None, bold=False, italic=False, style=None, space=5):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space)
    return p


def rich(parts, style=None, space=5):
    """parts = [(text, bold), ...] — for sentences with a bold term inside."""
    p = doc.add_paragraph(style=style)
    for text, bold in parts:
        r = p.add_run(text)
        r.bold = bold
    p.paragraph_format.space_after = Pt(space)
    return p


def bullets(items, style="List Bullet"):
    for it in items:
        if isinstance(it, list):
            parts = it
        elif isinstance(it, tuple):
            if len(it) == 2 and isinstance(it[1], str):
                parts = [(it[0], True), (it[1], False)]
            else:
                parts = [(it[0], bool(it[1]))]
        else:
            parts = [(it, False)]
        p = doc.add_paragraph(style=style)
        for text, bold in parts:
            p.add_run(text).bold = bold
        p.paragraph_format.space_after = Pt(2)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, htext in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(htext)
        run.bold = True
        run.font.size = Pt(9.5)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9.5)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ───────────────────────────────────────────────── read the deck back
from pptx import Presentation  # noqa: E402

prs = Presentation(str(PPTX))
SLIDE_TEXT = []
for s in prs.slides:
    lines = []
    for sh in s.shapes:
        if sh.has_text_frame:
            for l in sh.text_frame.text.split("\n"):
                if l.strip():
                    lines.append(l.strip())
    SLIDE_TEXT.append(lines)

src = MD.read_text(encoding="utf-8")
NOTES = {}
for m in re.finditer(r"^## (?:Slide (\d+)|Appendix (A\d)) — .*$", src, re.M):
    key = int(m.group(1)) if m.group(1) else m.group(2)
    note = re.search(r"^\*\*Speaker note(?: \([^)]*\))?\*\*\n(.*?)(?=^\*\*|^\n---|\Z)",
                     src[m.end():], re.M | re.S)
    if note:
        NOTES[key] = re.sub(r"\s+", " ", note.group(1)).strip().strip('"').replace("*", "").strip()

# ───────────────────────────────────────────────── cover
title = doc.add_paragraph()
r = title.add_run("Shotorko (সতর্ক)")
r.bold = True
r.font.size = Pt(26)
r.font.color.rgb = AMBER
sub = doc.add_paragraph()
r = sub.add_run("Presenter brief — the prototype, the code and the slides, in plain language")
r.font.size = Pt(13)
p = doc.add_paragraph()
r = p.add_run("GP FutureMakers 2026 · Concept Pitch Round · Saturday 17 October 2026, GP House · "
              "Live prototype https://abukawser.xyz")
r.font.size = Pt(9.5)
r.font.color.rgb = GREY

h1("If you only have five minutes")
bullets([
    "Shotorko is a Bangla phone-scam defender. It does not block the call. It interrupts the "
    "moment the scam asks for money, an OTP, a forwarding code or an install.",
    [("The one line that is the whole product: ", False), ("Stop the scam. Not the call.", True)],
    "Scams run on fear, urgency and secrecy. Every existing defence answers a different crime — "
    "spam filters answer volume, blocklists answer numbers, banks answer transfers. Nobody answers "
    "the three minutes inside the call.",
    "The buyer is not the victim. A parent loses the money; Grameenphone loses the trust. "
    "That is why an operator, not an app store, is the right home for this.",
    "Tomorrow you demonstrate three things only: Play the story on the live site, type a real scam "
    "line into নিজে যাচাই, and say the sentence 'the classifier reads intent, not a word list' while "
    "a reworded scam still pauses.",
    "If the venue network fails, the recorded video is your backup: Shotorko_Video_Pitch.mp4.",
])
para("The one sentence to open with and the one to close with — memorise both.", italic=True, color=GREY)
rich([("Open: ", True), ("\u201CStop the scam. Not the call.\u201D", False)])
rich([("Close: ", True), ("\u201Cসতর্ক থাকুন — stay alert, tell the family — and the Tk 30,000 stays "
                          "with the person who can least afford to lose it.\u201D", False)])

# ───────────────────────────────────────────────── the prototype
h1("1. The live prototype — what a visitor sees")
para("abukawser.xyz is the real prototype, not a mockup. No login, no install, it works on a phone, "
     "and the page says on every screen which parts are live and which are demo data.")
h2("How to drive it")
table(["Control", "What it does"], [
    ("▶ Play the story", "Runs the whole 19-step intervention unattended, about 46 seconds, with spoken "
                          "Bangla lines if the browser has a Bangla voice installed"),
    ("space bar", "Same play / stop toggle"),
    ("← →", "Previous / next scene — use these if the auto-run feels too slow"),
    ("⟲ Reset", "Back to a clean home screen. Press this before every demo"),
    ("🔊 Sound", "Turns the Bangla speech synthesis on and off — turn it OFF in a loud room"),
    ("On-device / Live agent", "Only appears when the backend answers /health. Switches the classifier "
                               "between the offline engine and the hosted model"),
], widths=[1.7, 5.3])
h2("The three phones on screen")
bullets([
    ("আম্মুর ফোন — Ammu's smartphone. The main device; it renders whichever scene you are in. "
     "This is where listening, the pause and the report happen.", False),
    ("আম্মুর পুরোনো বটন ফোন — the old button phone. It appears only on the কলের আগে scene, and that is "
     "the point: the pre-call warning arrives in the network, so it lands on a phone with no app, "
     "no store and no battery cost.", False),
    ("রাকিবের ফোন — the son's phone (আপনজন). Shows the family side: silence, then a red alert, then the "
     "one-tap call that ends the scam.", False),
    ("Right rail (presenter view) — four boxes: Scene, Engine output with the live risk meter, "
     "In production (what the real build would use), and What is simulated. Read from it if a judge "
     "asks 'is this real'.", False),
])
h2("The nine tabs, and what each one proves")
table(["Tab", "Screen", "What it proves", "Real or demo"], [
    ("0 হোম", "App home, four protection rows with switches", "Quiet when nothing is wrong; every "
     "feature is an opt-in toggle stored on the device", "UI real · stats demo"),
    ("1 কলের আগে", "Incoming call from 01711-234567 + the button phone", "A Bangla warning plays before "
     "the ring, on any handset, and the call is never blocked", "Reputation is demo · needs GP signalling"),
    ("2 কলের ভেতরে", "Live transcript, signal chips per line", "The phone understands meaning while the "
     "call runs; nothing is recorded, nothing leaves the handset", "Transcript scripted · classifier real"),
    ("3 থামুন", "The bKash send screen, then the overlay", "The intervention is on the action, not the "
     "call: one question, three buttons, the user keeps the last say", "Classifier real"),
    ("4 আপনজন", "QR, family secret word, alert state", "A child sets it up for a parent in two minutes; "
     "a cloned voice cannot pass the secret word", "QR is decorative · flow real"),
    ("5 এজেন্ট সতর্কতা", "Agent cash-out screen with a warning", "The shopkeeper becomes a checkpoint and "
     "asks the one question that saves people — advisory, never a block", "Feed is demo"),
    ("6 রিপোর্ট", "Speak the story → complaint letter", "A victim only speaks; the AI writes a filing-ready "
     "letter to GP fraud cell, the MFS provider and 999", "Writer real (local or hosted)"),
    ("7 স্ক্যাম রাডার", "Three scam networks, reports table", "Different numbers, same script — one network, "
     "not twenty separate complaints", "All demo data"),
    ("8 নিজে যাচাই", "Text box + situation chips + five reworded scripts", "The model itself, on any text a "
     "judge supplies, in Bangla or Banglish", "Real classifier"),
], widths=[0.95, 1.85, 2.9, 1.3])
rich([("The demo you should run if a judge doubts you: ", True),
      ("go to নিজে যাচাই and ask them for a sentence. Type it. It scores live. Then press one of the five "
       "\u201Csame intent, different words\u201D scripts and point at the target class versus what the model "
       "answered.", False)])

# ───────────────────────────────────────────────── the engine
h1("2. How the decision is actually made")
para("Both engines follow one rule: stop the action, not the call. A pause needs consent, a risk score "
     "at or above 50, and an actual request for money, a code, a forward or an install. Three conditions, "
     "AND-ed. That is why an ordinary family conversation about money stays silent.")
h2("The eight intent signals and their weights")
para("Weight is how much that signal contributes before normalisation. Highest first.", color=GREY, size=9.5)
table(["Signal", "Bangla label", "Weight", "Why it is weighted there"], [
    ("payment", "টাকা পাঠাতে বলা", "30", "The loss itself — money moving is the end goal of almost every scam"),
    ("credential", "গোপন কোড চাওয়া", "28", "An OTP, PIN, IVRF or a forwarding code hands over the account"),
    ("secrecy", "কাউকে বলতে নিষেধ", "26", "\u201CDon't tell anyone\u201D isolates the victim; no legitimate "
     "caller ever asks for this"),
    ("leverage", "আপনজনের বিপদ", "24", "A son, daughter or parent in trouble shuts judgement down"),
    ("install", "অ্যাপ বা স্ক্রিন শেয়ার", "24", "An unknown APK or screen share is total device capture"),
    ("authority", "ভুয়া ক্ষমতা", "22", "Osi, police, DB, CID, Rab, a case, an arrest — borrowed power"),
    ("urgency", "তাড়াতাড়ি করা হচ্ছে", "20", "\u201CRight now, before the line cuts\u201D prevents thinking"),
    ("reward", "লোভ দেখানো", "18", "Lottery, bonus, a job, a loan — the greed half, weaker than fear"),
], widths=[0.95, 1.5, 0.6, 3.9])
h2("The situation around the words")
para("The same sentence is harmless from a saved contact at noon and dangerous from a fresh unknown "
     "number mid-call. The engine scores the words, then the situation moves the number.", color=GREY, size=9.5)
table(["Modifier", "Adds", "Meaning"], [
    ("unknownCaller", "+12", "Nobody in the contacts book"),
    ("callActive", "+10", "The pressure is live, in the moment"),
    ("newReceiver", "+8", "Money is going somewhere it has never gone"),
    ("unusualAmount", "+8", "Bigger than this person normally moves"),
    ("lateNight", "+4", "Off hours, when family is asleep"),
    ("scare + (rush or isolate)", "+15", "The combination that defines a scam, not any single word"),
    ("isolate", "+6", "Secrecy on its own deserves extra weight"),
], widths=[2.0, 0.7, 4.3])
rich([("The arithmetic, in one line: ", True),
      ("sum the matched signal weights, add the situation modifiers, divide by 1.7, clamp to 0–100. "
       "If the result is 50 or more AND the caller asked for something AND the user opted in, থামুন fires.",
       False)])
h2("The eight scam types, and the rule that picks them")
para("Rules are ordered by specificity on purpose: the narrowest pattern must win, so \u201CbKash office, "
     "account suspended, give the OTP\u201D is read as an account takeover and never as a generic money "
     "request even though it mentions bKash.", color=GREY, size=9.5)
table(["Type", "Bangla label", "Fires when"], [
    ("forward", "কল ফরোয়ার্ডিং দখল", "a credential ask AND forwarding"),
    ("remote", "ফেক অ্যাপ / স্ক্রিন শেয়ার", "an install AND (credential OR payment)"),
    ("account", "অ্যাকাউন্ট ব্লক / ভুয়া MFS অফিসার", "authority AND OTP ask AND suspend"),
    ("refund", "ভুল পাঠানো / রিফান্ড", "refund AND an OTP ask"),
    ("police", "ভুয়া পুলিশ / মামলা", "authority AND (leverage OR payment OR urgency OR secrecy)"),
    ("emergency", "আপনজনের জরুরি বিপদ", "leverage AND (payment OR urgency)"),
    ("otp", "ওটিপি / অ্যাকাউন্ট হ্যাক", "OTP ask AND (reward OR suspend)"),
    ("prize", "ইনাম / চাকরি / ঋণ", "reward AND (payment OR install OR urgency)"),
], widths=[0.9, 2.3, 3.8])
h2("Worked example — the script on slide 2")
para("\u201Cহ্যালো, আমি থানার ওসি বলছি\u201D + \u201Cআপনার ছেলেকে ইয়াবাসহ ধরা হয়েছে\u201D + "
     "\u201Cএখনই ৩০ হাজার টাকা বিকাশ করেন\u201D + \u201Cকাউকে বলবেন না\u201D hits authority 22, "
     "leverage 24, payment 30, urgency 20 and secrecy 26 = 122. Situation (unknown caller, call active, "
     "new receiver, unusual amount) adds 38; scare with rush and isolation adds 15, secrecy adds 6. "
     "181 ÷ 1.7 = risk 100. It asks for money, so the pause fires.", size=10)
h2("Why rewording does not beat it")
bullets([
    "Matching is on phrase stems, not a fixed blocklist: \u201Ckauke bolben na\u201D, \u201Cgopon\u201D, "
    "\u201Cdon't tell\u201D all land on the same secrecy signal, in Bangla or Banglish.",
    "Digits are normalised (০১২৩৪৫৬৭৮৯ to 0123456789), punctuation and line breaks are flattened, so "
    "a dialect spelling still matches.",
    "The five ALT scripts in the console are the proof: bKash-office OTP, refund OTP, lottery fee, a "
    "pure-Banglish police demand, and a hospital emergency — five different wordings, same five intents, "
    "all still caught.",
])
rich([("Say this honestly if pushed: ", True),
      ("what runs in the browser today is a weighted stem matcher — a real classifier, but a fallback. "
       "Production is a fine-tuned Bangla model small enough for a cheap phone. The demo is built so the "
       "swap does not change the interface or the decision rule.", False)])

# ───────────────────────────────────────────────── backend
h1("3. The backend — one contract, two hosts")
para("The frontend only ever calls /functions/v1/app/* — health, analyze and complaint. "
     "vercel.json rewrites that to /api/app/*, so the exact same page runs on Vercel with Gemini or on "
     "Qoder Sites with Cloud Agents, with no frontend change.", size=10)
h2("The five safety properties, in plain words")
table(["Property", "What the code does", "Why it matters in the room"], [
    ("The model never speaks to the user", "Gemini returns signal KEYS only; every Bangla label comes from "
     "a server-side table", "A model cannot invent Bengali text, mistranslate, or leak a raw string into the "
     "page — no \u201C[object Object]\u201D on stage"),
    ("The server, not the model, decides", "wouldPause is accepted only if risk ≥ 50 AND a payment, "
     "credential or install signal is present", "\u201CThe model proposes, the server disposes\u201D — an "
     "over-eager answer cannot interrupt someone's call"),
    ("The key cannot be drained by a stranger", "15-minute cache on identical input, and a hard budget of "
     "60 upstream calls per minute; a cache hit costs nothing, a retry does", "One public page with a free "
     "API key is an open bill; this closes it"),
    ("Nothing from the provider is echoed", "a 429 becomes the single word {ok:false,error:\u201Crate_limited\u201D}; "
     "which quota window tripped is logged server-side only", "Quota text never reaches a visitor, and you "
     "still learn whether it was per-minute or per-day"),
    ("The page works with no backend at all", "if /health fails or a call errors, the on-device engine answers "
     "silently and shows why in a small fallback note", "A dead venue network cannot break the demo — "
     "it only makes it offline, which is the product's own design"),
], widths=[1.65, 2.65, 2.7])
h2("Model call mechanics (only if a technical judge asks)")
bullets([
    "gemini-2.5-flash by default, overridable with GEMINI_MODEL; a bad model id falls back to the default "
    "name once, and the model that actually answered is pinned until configuration changes.",
    "thinkingBudget is forced to 0 — a 2.5 model that thinks too much returns silence inside the output "
    "window. Requests then degrade in three attempts: strict JSON, plain JSON, then no format constraint.",
    "Deadlines are real: 18 seconds for analyze, 28 for complaint, with 500ms–20s per upstream attempt and "
    "a browser-side abort at 22s / 35s.",
    "Output is parsed by taking the first { to the last } and unwrapping code fences; anything malformed "
    "becomes agent_response_invalid, and a short body is treated as a miss rather than a wrong number.",
    "The complaint letter's recipient list is fixed in code — GP fraud cell, MFS provider, police / 999 — "
    "so the model cannot invent an authority to file with.",
])
h2("Two ways to run it")
table(["Host", "Backend", "Secret", "What to check"], [
    ("Vercel — abukawser.xyz", "api/app/[...path].ts, Edge runtime, Gemini", "GEMINI_API_KEY in the project "
     "environment; nothing else is required", "/functions/v1/app/health returns "
     "{ok:true,engine:\u201Cgemini\u201D,model:\u201Cgemini-2.5-flash\u201D} — verified today"),
    ("Qoder Sites", "functions/handler.ts on Cloud Agents", "QODER_PAT is a platform Secret — never in "
     "source, the bundle, a prompt or a log", "Same response contract, 24 checks pass offline"),
], widths=[1.5, 1.9, 2.4, 1.2])

# ───────────────────────────────────────────────── repo map
h1("4. The repository, file by file")
table(["File", "Lines", "What it is"], [
    ("web/index.html", "100", "One page, three phone frames, the tab strip, the transport bar, and the four "
     "footer boxes that state what is real, what is simulated and what rules the demo keeps"),
    ("web/assets/engine.js", "231", "The on-device classifier: 8 signals, 4 action triggers, 8 scam rules, "
     "the risk score, reputation lookup and the local complaint skeleton"),
    ("web/assets/app.js", "733", "Every screen, every button, the 19-step scripted story, the hybrid call to "
     "the backend and the automatic fall back to on-device"),
    ("web/assets/data.js", "86", "Demo only: the 6-turn scam call, 5 reworded scripts, the family, 16 reports, "
     "3 scam networks, 2 agent shops, the monthly stats"),
    ("web/assets/app.css", "—", "The dark theme, phone frames, switches, meters, print behaviour"),
    ("api/app/[...path].ts", "381", "Vercel Edge backend: Gemini, labels, the pause veto, cache, budget, "
     "timeouts, error codes"),
    ("functions/*.ts,.mjs", "372", "The same contract on Qoder Cloud Agents — handler, signals table, agent "
     "client"),
    ("vercel.json / .vercelignore", "—", "Static deploy of web/ with no build, plus the rewrite; dev/, pitch/ "
     "and functions/ never ship, and .env* is excluded"),
    ("pitch/deck-and-qa.md", "—", "Single source of truth for the deck: ten slides, three appendix, the "
     "sentences to speak, and the 23-question answer bank"),
    ("dev/build-pptx.py", "767", "Generates the .pptx — shapes, Bengali complex-script fonts through a:cs, "
     "real screenshots"),
    ("dev/build-deck.mjs", "—", "Generates pitch/deck.html from the markdown: 13 pages, arrow keys, prints to PDF"),
    ("dev/render-pptx.py", "—", "Redraws the .pptx with PIL to prove the layout, because PowerPoint is not "
     "on this machine"),
    ("dev/shoot.mjs / record.mjs", "—", "Capture the nine prototype screens and record the story video that "
     "slides and QR both point at"),
], widths=[1.8, 0.6, 4.6])
h2("The checks, and what each failure would mean")
table(["Command", "Checks", "Needs network?", "A failure means"], [
    ("node dev/engine-check.js", "33", "No", "A core demo claim broke — a scene that should pause went quiet "
     "or a family call would have been interrupted"),
    ("node dev/backend-check.mjs", "69", "No", "A server guarantee broke — labels, the pause veto, the budget, "
     "or an error leaking provider text"),
    ("node dev/function-check.mjs", "24", "No", "The Qoder Sites backend stopped matching the contract"),
    ("node dev/false-positive-battery.mjs local", "15 sentences", "No", "The slide 7 numbers on screen are no "
     "longer the measured numbers"),
    ("node dev/false-positive-battery.mjs <url>", "same sweep hosted", "Yes", "The hosted model disagrees with "
     "the offline engine — or Google's daily quota is throttling, which is not a failure"),
    ("node dev/live-probe.mjs [url]", "10", "Yes + chromium", "The deployed page genuinely does not work "
     "end to end"),
    ("python3 dev/render-pptx.py", "layout", "No", "Something on a slide is off-page, overflowing or colliding"),
], widths=[2.15, 0.85, 1.1, 2.9])
para("All six offline checks pass as of today: engine ALL CHECKS PASSED, backend 69/69, functions 24/24, "
     "battery benign 0/10 calm · 1/10 hostile · scams 4/5, deck layout clean over 216 text blocks, and "
     "https://abukawser.xyz/functions/v1/app/health answering with the live model.", size=10)

# ───────────────────────────────────────────────── numbers
h1("5. Every number you may be asked to defend")
table(["Figure", "Where", "Status", "Source / basis"], [
    ("191.09M mobile SIMs · 87.20M Grameenphone · 45.6% share", "4, 9", "Published", "BTRC subscriber data by "
     "operator, August 2026"),
    ("122.71M on 4G · 68M outside 4G · 69.18% mobile-internet penetration", "4, 9", "Published", "BTRC "
     "teledensity and penetration, August 2026 — 69.18% is mobile internet, not overall internet (77.89%)"),
    ("Tk 2.17T monthly MFS value · Tk 182.7B e-money balance · 258.7M registered accounts · 130.2M rural",
     "4, 12", "Published", "Bangladesh Bank MFS statistics, Tables 9 and 8 (the month is deliberately not "
     "claimed — Table 9 could not be verified as a July row from here)"),
    ("Over a billion MFS transactions in a month", "4", "Published", "Bangladesh Bank Table 9"),
    ("Tk 45.9B TAM · Tk 20.9B SAM · Tk 209M SOM", "9, 13", "Calculated", "subscribers × Tk 20 × 12, then 1% "
     "attach — arithmetic on the rows above"),
    ("Tk 105M / 209M / 419M / 628M at 0.5 / 1 / 2 / 3%", "12", "Calculated", "Same arithmetic, gross revenue, "
     "not profit"),
    ("4M+ premium subscribers · 0.75% average conversion", "8, 9, 13", "Published", "Truecaller AB, "
     "10 March 2026, Q4 2025 — global, not Bangladeshi"),
    ("USD 3.99 a month", "8", "Published", "Verizon Call Filter Plus pricing"),
    ("USD 3.5B reported lost to imposter scams in 2025", "A1", "Published", "US FTC, 15 June 2026 — cited as a "
     "shape, not a size for Bangladesh"),
    ("Tk 20 a month · 1% attach", "8, 9, 12", "OUR ASSUMPTION", "Proposed price and base case; not a surveyed "
     "willingness to pay — say so first, every time"),
    ("risk 100/100 · risk 95/100 · 0 of 10 · 4 of 5 · 69 checks", "6, 7", "Prototype result", "Local battery "
     "and the live console — preliminary, 15 test sentences, not a field trial"),
    ("No Bangladeshi national fraud-loss figure", "A1", "Deliberately not claimed", "It does not exist in "
     "published form, and we refused to invent one — this is a credibility point, use it out loud"),
], widths=[2.45, 0.6, 1.1, 2.85])

# ───────────────────────────────────────────────── slides
h1("6. The slides, one by one")
para("For each slide: exactly what is on it (read out of the .pptx), the sentence to say "
     "(read out of the markdown), and the trap to avoid.", color=GREY, size=9.5, italic=True)

TRAPS = {
    1: ("Do not explain the name twice — the logo is on screen. Do not start with the technology; "
        "start with the fear."),
    2: ("This is an illustrative scenario, not a documented incident — if asked, say so plainly. "
        "Do not race through the timeline; the 3:00 row is the whole pitch."),
    3: ("Never say existing defences are useless. Say they answer a different crime — carrier filters are "
        "good at what they were built for."),
    4: ("Do not say \u201CJuly 2026\u201D for the Bangladesh Bank rows; the tables are named, the month is "
        "not claimed. Say 258.7M registered accounts, not 258.7M people."),
    5: ("Do not explain all four moments at once. Two need no smartphone — that is the line to land."),
    6: ("Do not read the three buttons in an order that sounds like a flowchart. Say: verify, family, or "
        "the user's own call — and the last one is respected."),
    7: ("If a judge asks whether this is an LLM: today's browser engine is a weighted stem classifier, and "
        "say exactly that. The 95/100 screenshot is a hosted model call — demo evidence, not field accuracy."),
    8: ("Do not claim Truecaller proves Bangladeshis will pay Tk 20. It proves people pay for caller safety. "
        "Also say gross revenue, not profit."),
    9: ("Redo the arithmetic out loud if challenged: 191 million times 20 times 12. Never present TAM as "
        "something you will capture."),
    10: ("Land the four asks separately: one is data, one is a network slot, one is a gating test GP must "
        "validate, one is a district. End on the site, not on yourself."),
    "A1": ("Only open this if a number is challenged. Point at the calculation row and admit the assumptions "
           "before the judge finds them."),
    "A2": ("This is a sensitivity page, not a forecast. Say the only assumption that moves it is attach rate."),
    "A3": ("Reference only — never presented. Navigate here in the Q&A to show the reframed question, then "
           "answer it."),
}
CRITERION = {
    1: "Clarity of presentation", 2: "Depth of problem understanding", 3: "Depth of problem understanding · "
    "differentiation", 4: "Depth of problem understanding · impact", 5: "Solution strength & differentiation",
    6: "Solution strength · AI application", 7: "AI application & technical rationale",
    8: "Feasibility & scalability · business sense", 9: "Feasibility & scalability · impact",
    10: "Impact · clarity of the ask", "A1": "Credibility under challenge",
    "A2": "Feasibility · unit economics", "A3": "Q&A handling",
}
KEY = {1: "what works today, what GP must enable, and what the pilot has to prove"}
FOOTER = "Shotorko · সতর্ক · GP FutureMakers 2026"
PAGE = re.compile(r"(?:[0-9]{1,2} / 10|[Aa][0-9] / A3)")
for i, raw in enumerate(SLIDE_TEXT, 1):
    lines = [l for l in raw if l != FOOTER and not PAGE.fullmatch(l)]
    key = i if i <= 10 else "A%d" % (i - 10)
    label = f"Slide {key}" if isinstance(key, int) else f"Appendix {key}"
    h2(f"{label} — " + (" · ".join(lines[:2]) if i > 10 else lines[0]))
    body = lines[2:] if i > 10 else lines[1:]
    body = [b for b in body if "GP FutureMakers 2026 · Concept" not in b]
    para("On screen: " + " | ".join(body)[:900], size=9.5, color=GREY)
    if key in NOTES:
        rich([("Say: ", True), (NOTES[key], False)], space=3)
    rich([("Scores: ", True), (CRITERION.get(key, ""), False)], space=3)
    rich([("Trap: ", True), (TRAPS.get(key, ""), False)], space=3)

# ───────────────────────────────────────────────── Q&A
h1("7. The questions that will actually come")
para("Answer in three beats: the direct answer, the honest limit, and what would change the answer. "
     "Never quote the demo dashboard as if it were a field trial.", color=GREY, size=9.5, italic=True)
QA = [
    ("How big is the market?",
     "191 million mobile subscribers; at Tk 20 a month that is Tk 45.9B a year across all operators and "
     "Tk 20.9B on Grameenphone's 87.2 million base. That is gross revenue potential before cost, revenue "
     "share and tax. Year one we only need 1% of GP's base — Tk 209M — and that 1% is our assumption, not "
     "a survey result."),
    ("Will people really pay Tk 20?",
     "Evidence that users pay for communication safety exists: Truecaller passed 4 million paying "
     "subscribers with a 0.75% average conversion in Q4 2025, and Verizon charges USD 3.99 for Call Filter "
     "Plus. Both are global or US, so they show willingness, not our price. Tk 20 rides on a bill the "
     "subscriber already pays, which is why the operator is the right channel. The pilot has to measure "
     "real willingness to pay."),
    ("Why not just an app?",
     "68 million subscribers are outside 4G and many of the people who get these calls use a button phone "
     "costing around Tk 1,000. An app can only protect the person with the right handset, the store "
     "account and the data plan. Two of the four moments — the warning before the ring and the divert "
     "gate — happen in the network, so they reach everyone. The on-device half needs a smartphone and is "
     "offered as the upgrade path."),
    ("How does it listen on a phone with no apps?",
     "It does not listen on that phone. The pre-call warning is a network announcement played before the "
     "call connects, from reputation the operator already holds. Speech understanding happens on a "
     "smartphone, on the handset, only for users who opt in. Be precise about which component does which "
     "job — this is the question that separates a real plan from a slogan."),
    ("Isn't this mass surveillance of calls?",
     "No, and the design refuses it: opt-in per feature, processing on the handset, no recording, no audio "
     "uploaded, nothing stored server-side, and the user keeps the last button on every warning. The only "
     "thing that leaves the phone is the complaint the victim chooses to file."),
    ("What is your false-positive rate?",
     "Preliminary and small: ten ordinary sentences paused zero times in a friendly context and once under "
     "every hostile flag forced on — and that one was \u201Ctell abbu nothing, I need 3000 taka\u201D, which "
     "under unknown caller, mid-call and an unusual amount is the actual attack. Four of five reworded scam "
     "scripts paused; the fifth was a bare forwarding instruction with no money, no code and no pressure, "
     "and it stayed quiet on purpose because that case is stopped network-side. Fifteen sentences is not a "
     "field trial, which is exactly what the 90-day pilot measures, and every warning ships with a "
     "\u201Cঠিক ছিল\u201D button so the false-alarm count is a live number."),
    ("Have you tested regional Bangla and dialects?",
     "Bangla, Banglish and dialect spellings are normalised and handled in text, and Banglish inputs are "
     "in the test set. Real speech across Sylheti, Chittagonian and no accent has not been tested on "
     "recorded audio — no speech front-end exists yet. Four weeks of real Bangla speech that cannot hold "
     "the false-positive rate is the condition that would change our approach."),
    ("What does it cost you to run?",
     "On-device inference means no cloud call per warning, and there is no app-store royalty because the "
     "product is a network service. What still costs money is not yet measured: USSD and airtime, complaint "
     "handling, model updates and support. That is an honest unknown, not a solved number."),
    ("How does Grameenphone profit?",
     "Revenue share on the pack is the visible half. The larger half is avoided cost and avoided damage: "
     "fraud refunds, hotline minutes, churn after a publicised loss, and the 182.7B taka of e-money float "
     "that only works if people trust the network. Whether GP nets more than it spends is precisely what "
     "the pilot is designed to show — nobody should believe that number on a slide."),
    ("What do you need from us, today?",
     "Four things: anonymised fraud-call audio and CDR patterns for training, one network announcement "
     "slot for the pre-call warning, a divert-gating test on the codes Grameenphone already publishes — "
     "##002# to cancel all forwarding and *#21# to read its status, subject to your validation — and one "
     "pilot district for 90 days with agreed success criteria."),
    ("Who is the team?",
     "Answer as the truth. One builder has shipped a working Bangla prototype, a backend with safety "
     "guarantees and 126 passing checks. What is missing is a team, and the pilot is how you recruit for "
     "the right problem. Do not overstate."),
    ("What would make you stop?",
     "Real Bangla speech that cannot hold the false-positive rate at a level people will tolerate, or a "
     "pilot where opt-in consent does not convert. A defender that cries wolf is worse than no defender, "
     "because it teaches people to ignore warnings."),
    ("Why now?",
     "Bangladesh already moved its money onto rails worth over a trillion taka a month, and rural accounts "
     "now outnumber urban ones. The fraud followed the money. The defence did not."),
    ("Isn't bKash or Nagad already doing this?",
     "They see the transfer, never the call. By the time the transaction exists, the victim has already "
     "been convinced. The intervention has to be upstream of the transfer, which is why a network-side "
     "signal beats a wallet-side rule."),
]
for q, a in QA:
    p = doc.add_paragraph()
    r = p.add_run("Q — " + q)
    r.bold = True
    p.paragraph_format.space_after = Pt(1)
    para(a, size=10, space=7)

# ───────────────────────────────────────────────── glossary
h1("8. Bangla on the screen — pronunciation and meaning")
table(["On screen", "Says", "Meaning"], [
    ("সতর্ক", "shotor-ko", "Be alert / cautious — the product name"),
    ("থামুন", "tha-mun", "Stop / pause — the intervention screen"),
    ("আপনজন", "a-pon-jon", "One's own people — the family circle feature"),
    ("নিজে যাচাই", "ni-je ja-chai", "Check it yourself — the free-play console"),
    ("ঠিক ছিল", "thik chilo", "That was fine — the false-alarm feedback button"),
    ("কাউকে বলবেন না", "kauke bolben na", "Don't tell anyone — the secrecy signal"),
    ("থানার ওসি", "thana-r osi", "The police station officer-in-charge — the classic fake authority"),
    ("ইয়াবা", "yaba", "Yaba tablets — the drug in the fake arrest script"),
    ("আম্মু / রাকিব", "ammu / rakib", "Mother / the son's name in the demo"),
    ("শেখেরচালা", "shekher-chaala", "The demo family secret word — an unusual village name, easy to verify "
     "aloud"),
    ("সব ঠিক আছে", "shob thik aachhe", "Everything is fine — the home screen when nothing is wrong"),
], widths=[1.5, 1.5, 3.9])

# ───────────────────────────────────────────────── known issues
h1("9. Known weak points — so nothing surprises you")
bullets([
    ("The live prototype still says \u201CTk 800 button phone\u201D in the pre-call screen copy, while the "
     "slides now say about Tk 1,000. A judge reading both will notice. One line in "
     "web/assets/app.js and a redeploy fixes it — say the words on the slides, not the ones on the page."),
    ("Google's free Gemini quota may already be exhausted tomorrow. If it is, /health still answers, the "
     "Live agent toggle still appears, and every call falls back to the on-device engine with a small "
     "fallback note. The deck's numbers are the on-device ones precisely because they re-run identically "
     "anywhere. Do not present a throttled result as a classifier result."),
    ("Venue Wi-Fi is the other risk. Play the story once on the venue network before you are on stage. "
     "If it fails, the recorded video (Shotorko_Video_Pitch.mp4) shows the same intervention with no "
     "network at all, and the on-device engine works offline by design."),
    ("The QR on slide 10 has only ever been scanned from this machine. Scan it once from a phone on the "
     "venue network."),
    ("The nine screenshots in the deck are captures of the current build. If you change the prototype "
     "tonight, re-run node dev/shoot.mjs && python3 dev/build-pptx.py or the slides will show an older UI."),
    ("Re-check BTRC on the morning of 17 October — the subscriber tables are re-published monthly, and a "
     "judge who works at an operator will know."),
    ("Nothing in the repo contains a secret: .env* is git-ignored and excluded from the deploy, and "
     "QODER_PAT is a platform Secret only. Keep it that way if you push tonight."),
])

# ───────────────────────────────────────────────── checklist
h1("10. Pre-flight, in order")
table(["#", "Do this", "Time"], [
    ("1", "python3 dev/render-pptx.py on the submitted .pptx — the last line must read "
          "\u201Cclean over 216 text blocks\u201D", "2 min"),
    ("2", "Open the .pptx and flip through all 13 pages once; check the fonts render Bangla on your machine",
     "3 min"),
    ("3", "node dev/engine-check.js && node dev/backend-check.mjs — both must pass", "1 min"),
    ("4", "Open abukawser.xyz on your phone, press ⟲ Reset then ▶ Play the story, end to end", "2 min"),
    ("5", "Type a scam line you were actually sent into নিজে যাচাই and read the result aloud — this is the "
          "demo you will repeat on stage", "2 min"),
    ("6", "Scan the QR from a phone on the venue network", "30 sec"),
    ("7", "Confirm Shotorko_Video_Pitch.mp4 plays on the presentation laptop, muted", "1 min"),
    ("8", "Say the open line and the close line out loud, then the four asks on slide 10 in order", "3 min"),
    ("9", "Re-read section 9 above — the honest limits are your strongest material, not your weakest", "5 min"),
    ("10", "The deck is already submitted — treat that .pptx as frozen. If you change anything tonight, "
     "run python3 dev/build-brief.py after it, so this document still matches what is on screen", "1 min"),
], widths=[0.35, 5.6, 1.0])

para("Generated by dev/build-brief.py from pitch/Shotorko-GP-FutureMakers-2026.pptx and "
     "pitch/deck-and-qa.md, so it matches the deck that was actually built.",
     size=8.5, color=GREY, italic=True)

doc.save(str(OUT))
print(f"wrote {OUT.relative_to(ROOT)} · {len(doc.paragraphs)} paragraphs · "
      f"{len(doc.tables)} tables")
